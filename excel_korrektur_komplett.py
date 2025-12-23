#!/usr/bin/env python3
"""
Komplette Korrektur der Theken-Abrechnungsliste 2026
Führt alle 4 Schritte automatisch durch
"""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from datetime import datetime, timedelta
import calendar

DATEI_PFAD = '/Users/dr.rolandspiehl/skills-introduction-to-github2/Theken-Abrechnungsliste2026.xlsx'

def hole_wochenendtage(jahr, monat):
    """Gibt alle Freitage, Samstage und Sonntage eines Monats zurück"""
    wochenendtage = []
    tage_im_monat = calendar.monthrange(jahr, monat)[1]

    for tag in range(1, tage_im_monat + 1):
        datum = datetime(jahr, monat, tag)
        # weekday(): 4=Fr, 5=Sa, 6=So
        if datum.weekday() in [4, 5, 6]:
            wochentag = ['Mo', 'Di', 'Mi', 'Do', 'Fr', 'Sa', 'So'][datum.weekday()]
            wochenendtage.append((datum, wochentag))

    return wochenendtage

def main():
    print("=" * 70)
    print("THEKEN-ABRECHNUNGSLISTE 2026 - AUTOMATISCHE KORREKTUR")
    print("=" * 70)

    # Lade Arbeitsbuch
    print(f"\n📁 Lade: {DATEI_PFAD}")
    wb = openpyxl.load_workbook(DATEI_PFAD)
    print(f"✓ Geladen. Blätter: {wb.sheetnames}")

    # Monatsnamen mapping
    monat_namen = {
        'Jan': 1, 'Feb': 2, 'Mrz': 3, 'Apr': 4, 'Mai': 5, 'Jun': 6,
        'Jul': 7, 'AUG': 8, 'Sep': 9, 'Okt': 10, 'Nov': 11, 'Dez': 12
    }

    # Basis-Zeilennummer
    basis_zeilen_nr = 20001

    print("\n" + "=" * 70)
    print("VERARBEITE ALLE MONATE...")
    print("=" * 70)

    for sheet_name in wb.sheetnames:
        if sheet_name not in monat_namen:
            print(f"⚠️  Überspringe unbekanntes Blatt: {sheet_name}")
            continue

        monat_nr = monat_namen[sheet_name]
        ws = wb[sheet_name]

        print(f"\n📊 {sheet_name} (Monat {monat_nr}/12)")

        # SCHRITT 1 & 4: Hole alle Wochenendtage für diesen Monat
        wochenendtage = hole_wochenendtage(2026, monat_nr)
        print(f"   ✓ {len(wochenendtage)} Wochenendtage gefunden")

        # Berechne Start-Zeilennummer für diesen Monat
        start_zeilen_nr = basis_zeilen_nr + (monat_nr - 1) * 30

        # SCHRITT 2 & 3: Lösche alte Daten (ab Zeile 2) und erstelle neu
        # Lösche alle Zeilen außer Header
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

        # SCHRITT 2 & 3: Erstelle genau 30 Datenzeilen
        for zeile_idx in range(30):
            zeile_nr = zeile_idx + 2  # Zeile 2-31 (Header ist Zeile 1)
            zeilen_id = start_zeilen_nr + zeile_idx

            # Schreibe Zeilennummer
            ws.cell(zeile_nr, 1, zeilen_id)

            # Wenn es einen Wochenendtag gibt, fülle die Daten ein
            if zeile_idx < len(wochenendtage):
                datum, wochentag = wochenendtage[zeile_idx]
                ws.cell(zeile_nr, 2, datum)  # Datum
                ws.cell(zeile_nr, 3, wochentag)  # Wochentag
            # Sonst bleiben Spalten B, C leer

            # Spalten D-G (Name, Vorname, von, bis) bleiben leer

        # SCHRITT 3: DIN-A4 Formatierung
        # Seiteneinrichtung
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToHeight = 1
        ws.page_setup.fitToWidth = 1

        # Ränder
        ws.page_margins.left = 0.7
        ws.page_margins.right = 0.7
        ws.page_margins.top = 0.75
        ws.page_margins.bottom = 0.75

        # Druckbereich
        ws.print_area = f'A1:G31'

        # Spaltenbreiten anpassen
        ws.column_dimensions['A'].width = 8   # #
        ws.column_dimensions['B'].width = 12  # Datum
        ws.column_dimensions['C'].width = 5   # WT
        ws.column_dimensions['D'].width = 15  # Name
        ws.column_dimensions['E'].width = 15  # Vorname
        ws.column_dimensions['F'].width = 8   # von
        ws.column_dimensions['G'].width = 8   # bis

        print(f"   ✓ 30 Zeilen erstellt (IDs: {start_zeilen_nr}-{start_zeilen_nr+29})")
        print(f"   ✓ DIN-A4 Formatierung gesetzt")

    # Backup erstellen
    backup_pfad = DATEI_PFAD.replace('.xlsx', '_BACKUP.xlsx')
    print(f"\n💾 Erstelle Backup: {backup_pfad}")

    try:
        import shutil
        shutil.copy2(DATEI_PFAD, backup_pfad)
        print(f"✓ Backup erstellt")
    except Exception as e:
        print(f"⚠️  Backup-Warnung: {e}")

    # Speichern
    print(f"\n💾 Speichere: {DATEI_PFAD}")
    wb.save(DATEI_PFAD)
    print(f"✓ Gespeichert!")

    # Zusammenfassung
    print("\n" + "=" * 70)
    print("ZUSAMMENFASSUNG")
    print("=" * 70)
    print("✓ Schritt 1: Datumsangaben korrigiert (2026)")
    print("✓ Schritt 2: Zeilennummern 20001-20360 (30 pro Monat)")
    print("✓ Schritt 3: DIN-A4 Formatierung gesetzt")
    print("✓ Schritt 4: Alle Wochenendtage eingetragen")
    print("\n🎉 FERTIG!")
    print("=" * 70)

    # Detaillierte Wochenendtage-Übersicht
    print("\n📅 WOCHENENDTAGE 2026 - ÜBERSICHT:")
    print("-" * 70)

    for sheet_name in wb.sheetnames:
        if sheet_name in monat_namen:
            monat_nr = monat_namen[sheet_name]
            wochenendtage = hole_wochenendtage(2026, monat_nr)
            print(f"\n{sheet_name} 2026: {len(wochenendtage)} Tage")
            for datum, wt in wochenendtage:
                print(f"  {wt} {datum.strftime('%d.%m.%Y')}")

if __name__ == "__main__":
    main()
